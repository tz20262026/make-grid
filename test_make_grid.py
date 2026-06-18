from pathlib import Path
from docx import Document
from docx.document import Document as DocumentClass

import make_grid


def test_build_grid_returns_document():
    doc = make_grid.build_grid_document()
    assert isinstance(doc, DocumentClass)


def test_grid_has_correct_dimensions():
    doc = make_grid.build_grid_document()
    table = doc.tables[0]
    assert len(table.rows) == make_grid.ROWS
    assert len(table.columns) == make_grid.COLS


def test_cell_numbers_are_sequential():
    doc = make_grid.build_grid_document()
    table = doc.tables[0]
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            expected = str(i * make_grid.COLS + j + 1)
            assert cell.paragraphs[0].runs[0].text == expected


def test_total_cells():
    doc = make_grid.build_grid_document()
    table = doc.tables[0]
    count = sum(1 for row in table.rows for _ in row.cells)
    assert count == make_grid.ROWS * make_grid.COLS


def test_save_creates_valid_docx(tmp_path):
    doc = make_grid.build_grid_document()
    out = tmp_path / "test_grid.docx"
    doc.save(out)
    assert out.exists()
    loaded = Document(out)
    assert len(loaded.tables) == 1
    assert len(loaded.tables[0].rows) == make_grid.ROWS


def test_emu_per_twip_constant():
    # 1 twip = 914400 EMU/inch ÷ 1440 twips/inch = 635 EMU
    assert make_grid.EMU_PER_TWIP == 635


def test_output_path_is_relative_to_script():
    assert make_grid.OUTPUT_PATH.parent == Path(make_grid.__file__).parent
