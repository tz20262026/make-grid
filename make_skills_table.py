from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# タイトル
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Claude Code スキル一覧")
run.font.size = Pt(18)
run.font.bold = True

doc.add_paragraph()

# 表データ
rows_data = [
    ("コマンド", "機能", ""),
    ("/init",   "新プロジェクト一括セットアップ", "CLAUDE.md・git・README・.gitignore を自動生成"),
    ("/review", "コード差分の自動レビュー＆修正",  "バグ・セキュリティ・品質を自律チェックして即修正"),
    ("/deploy", "テスト→ビルド→lint→デプロイ",    "コミットからVercel本番デプロイまで一気通貫"),
    ("/status", "全プロジェクトのgit状態一覧",     "デスクトップ全プロジェクトの未コミット・差分を表示"),
]

table = doc.add_table(rows=len(rows_data), cols=3)
table.style = 'Table Grid'

col_widths = [Cm(3.5), Cm(6.5), Cm(8.5)]
header_color = RGBColor(0x2C, 0x3E, 0x50)
header_text_color = RGBColor(0xFF, 0xFF, 0xFF)
alt_color = RGBColor(0xF2, 0xF2, 0xF2)

for i, (cmd, func, detail) in enumerate(rows_data):
    row = table.rows[i]
    cells = row.cells
    texts = [cmd, func, detail]

    for j, (cell, text) in enumerate(zip(cells, texts)):
        cell.width = col_widths[j]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        run.font.size = Pt(11 if i > 0 else 12)
        run.font.bold = (i == 0)

        if i == 0:
            run.font.color.rgb = header_text_color
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), '2C3E50')
            tcPr.append(shd)
        elif i % 2 == 0:
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'F2F2F2')
            tcPr.append(shd)

doc.add_paragraph()
note = doc.add_paragraph("※ どのプロジェクトフォルダでも使用できます。Claude Code で /コマンド名 と入力するだけで実行されます。")
note.runs[0].font.size = Pt(9)
note.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

out = Path(__file__).parent / "claude_skills_table.docx"
doc.save(out)
print(f"saved: {out}")
