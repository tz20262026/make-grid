from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

COLS = 5
ROWS = 4
PAGE_WIDTH_CM = 29.7
PAGE_HEIGHT_CM = 21.0
MARGIN_CM = 1.0
CELL_NUMBER_FONT_SIZE = 8
CELL_MARGIN_DXA = 80
# 1 twip = 635 EMU  (1/20 pt × 12700 EMU/pt)
EMU_PER_TWIP = 635

OUTPUT_PATH = Path(__file__).parent / f"grid_{COLS * ROWS}.docx"


def _set_cell_margins(cell, margin_dxa: int) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for side in ('top', 'left', 'bottom', 'right'):
        m = OxmlElement(f'w:{side}')
        m.set(qn('w:w'), str(margin_dxa))
        m.set(qn('w:type'), 'dxa')
        tcMar.append(m)
    tcPr.append(tcMar)


def _set_row_height(row, height_cm: float) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(int(Cm(height_cm).emu / EMU_PER_TWIP)))
    trHeight.set(qn('w:hRule'), 'exact')
    trPr.append(trHeight)


def build_grid_document() -> Document:
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    section.orientation = 1  # WD_ORIENT.LANDSCAPE
    section.top_margin = Cm(MARGIN_CM)
    section.bottom_margin = Cm(MARGIN_CM)
    section.left_margin = Cm(MARGIN_CM)
    section.right_margin = Cm(MARGIN_CM)

    usable_width_cm = PAGE_WIDTH_CM - 2 * MARGIN_CM
    usable_height_cm = PAGE_HEIGHT_CM - 2 * MARGIN_CM
    col_width = Cm(usable_width_cm / COLS)
    row_height_cm = usable_height_cm / ROWS

    table = doc.add_table(rows=ROWS, cols=COLS)
    table.style = 'Table Grid'

    for col in table.columns:
        for cell in col.cells:
            cell.width = col_width

    for i, row in enumerate(table.rows):
        _set_row_height(row, row_height_cm)
        for j, cell in enumerate(row.cells):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(i * COLS + j + 1))
            run.font.size = Pt(CELL_NUMBER_FONT_SIZE)
            run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            _set_cell_margins(cell, CELL_MARGIN_DXA)

    return doc


def main() -> None:
    doc = build_grid_document()
    doc.save(OUTPUT_PATH)
    print(f"saved: {OUTPUT_PATH}")


if __name__ == '__main__':
    main()
